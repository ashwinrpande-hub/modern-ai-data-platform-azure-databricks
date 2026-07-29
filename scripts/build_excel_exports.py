#!/usr/bin/env python3
"""Build one Excel workbook per medallion layer from extract_data_model.py's output.
Each workbook: an INDEX summary sheet + one sheet per table with its 500-row sample.
Pure pandas/openpyxl -- no Spark session needed, run after extract_data_model.py."""
import json, os, re
import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

IN_DIR = os.path.join("data", "exports", "_extraction")
OUT_DIR = os.path.join("data", "exports", "excel")

LAYER_LABEL = {
    "acme_bronze": "Bronze — raw, source-shaped",
    "acme_silver": "Silver — 3NF, insert-only, hash-keyed",
    "acme_gold": "Gold — star schema + DV2.0 + ML/AI",
    "acme_products": "Products — governed output ports (views)",
}
LAYER_COLOR = {
    "acme_bronze": "CD7F32", "acme_silver": "B0B7BD",
    "acme_gold": "D4AF37", "acme_products": "4472C4",
}


def sheet_name(schema, table):
    raw = re.sub(r"[\[\]\:\*\?/\\]", "_", f"{schema}.{table}")
    return raw[:31]


def main():
    with open(os.path.join(IN_DIR, "extraction_result.json"), encoding="utf-8") as f:
        data = json.load(f)

    os.makedirs(OUT_DIR, exist_ok=True)
    by_catalog = {}
    for t in data["tables"]:
        by_catalog.setdefault(t["table_catalog"], []).append(t)

    for catalog, tables in sorted(by_catalog.items()):
        tables = sorted(tables, key=lambda x: (x["table_schema"], x["table_name"]))
        xlsx_path = os.path.join(OUT_DIR, f"{catalog}.xlsx")
        with pd.ExcelWriter(xlsx_path, engine="openpyxl") as xw:
            idx_df = pd.DataFrame([{
                "schema": t["table_schema"], "table": t["table_name"], "type": t["table_type"],
                "row_count_live": t.get("row_count"), "rows_sampled_here": t.get("sample_rows", 0),
                "columns": len(t.get("columns", [])), "comment": t.get("comment") or "",
            } for t in tables])
            idx_df.to_excel(xw, sheet_name="INDEX", index=False)
            for t in tables:
                csv_path = t.get("sample_csv")
                sname = sheet_name(t["table_schema"], t["table_name"])
                df = pd.read_csv(csv_path) if csv_path and os.path.exists(csv_path) \
                    else pd.DataFrame({"note": [f"no sample available: {t.get('sample_error','')}"]})
                df.to_excel(xw, sheet_name=sname, index=False)

        wb = load_workbook(xlsx_path)
        color = LAYER_COLOR[catalog]
        header_fill = PatternFill("solid", fgColor=color)
        header_font = Font(bold=True, color="FFFFFF" if catalog != "acme_gold" else "000000")
        for ws in wb.worksheets:
            for cell in ws[1]:
                cell.fill, cell.font = header_fill, header_font
                cell.alignment = Alignment(vertical="center")
            ws.freeze_panes = "A2"
            for col_cells in ws.columns:
                length = max((len(str(c.value)) if c.value is not None else 0) for c in col_cells)
                ws.column_dimensions[get_column_letter(col_cells[0].column)].width = min(max(length + 2, 10), 42)
        ws0 = wb["INDEX"]
        ws0.insert_rows(1)
        ws0["A1"] = f"{catalog}  —  {LAYER_LABEL[catalog]}"
        ws0["A1"].font = Font(bold=True, size=13)
        wb.move_sheet("INDEX", offset=-len(wb.sheetnames))
        wb.save(xlsx_path)
        print(f"wrote {xlsx_path}  ({len(tables)} tables)")


if __name__ == "__main__":
    main()
